"""PDF (ReportLab) and Excel (openpyxl) report generation."""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

HEADER_FILL = PatternFill("solid", fgColor="1D4ED8")
HEADER_FONT = Font(color="FFFFFF", bold=True)

styles = getSampleStyleSheet()
TITLE = styles["Title"]
BODY = styles["BodyText"]


def _excel_sheet(ws, headers: list[str], data: list[list], title: str | None = None) -> None:
    row = 1
    if title:
        ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
        row = 3
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    for record in data:
        row += 1
        for col, value in enumerate(record, start=1):
            ws.cell(row=row, column=col, value=value)
    for col, h in enumerate(headers, start=1):
        width = max(len(str(h)), *(len(str(r[col - 1])) for r in data)) if data else len(str(h))
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = min(width + 3, 40)


def class_ranking_xlsx(rows: list[dict], class_name: str, term: str, year: int) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Ranking"
    headers = ["Rank", "Admission No", "Name", "Class", "Average", "Mean Grade",
               "Points", "Subjects", "Pass Rate %"]
    data = [[r["rank"], r["admission_no"], r["name"], r["class_name"], r["average"],
             r["mean_grade"], r["total_points"], r["subjects"], r["pass_rate"]] for r in rows]
    _excel_sheet(ws, headers, data, f"Class Ranking - {class_name} ({term} {year})")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def subject_analysis_xlsx(df, term: str, year: int) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Subject Analysis"
    headers = list(df.columns)
    data = [list(r) for r in df.itertuples(index=False, name=None)]
    _excel_sheet(ws, headers, data, f"Subject Analysis ({term} {year})")
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def report_card_pdf(report: dict) -> bytes:
    """Individual student report card."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=18 * mm, bottomMargin=18 * mm,
                            title=f"Report Card - {report['student']['name']}")
    student = report["student"]
    story = [
        Paragraph("Chokur System", TITLE),
        Paragraph("Termly Report Card", BODY),
        Spacer(1, 6 * mm),
        Paragraph(
            f"<b>Name:</b> {student['name']} &nbsp;&nbsp; "
            f"<b>Admission No:</b> {student['admission_no']}<br/>"
            f"<b>Class:</b> {student['class_name']} &nbsp;&nbsp; "
            f"<b>Term:</b> {report['term']} {report['year']}",
            BODY,
        ),
        Spacer(1, 6 * mm),
    ]

    table_data = [["Subject", "Code", "Mark", "Grade", "Points"]]
    for s in report["subjects"]:
        table_data.append([s["subject"], s["code"], f"{s['score']:.0f}",
                           s["grade"], str(s["points"])])
    table = Table(table_data, hAlign="LEFT", colWidths=[60 * mm, 20 * mm, 25 * mm,
                                                        25 * mm, 25 * mm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1D4ED8")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
        ("ALIGN", (2, 1), (-1, -1), "CENTER"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    story.append(table)
    story.append(Spacer(1, 6 * mm))
    story.append(Paragraph(
        f"<b>Aggregate Points:</b> {report['total_points']} &nbsp;&nbsp; "
        f"<b>Average:</b> {report['average']:.1f}% &nbsp;&nbsp; "
        f"<b>Mean Grade:</b> {report['mean_grade']}<br/>"
        f"<b>Class Position:</b> {report['rank']} of {report['class_size']} &nbsp;&nbsp; "
        f"<b>Result:</b> {'PASS (C+ and above)' if report['passed'] else 'BELOW C+'}",
        BODY,
    ))
    story.append(Spacer(1, 8 * mm))
    story.append(Paragraph(
        "Teacher's remarks: ____________________________________________<br/><br/>"
        "Signature: ____________________ &nbsp; Date: ______________",
        BODY,
    ))
    doc.build(story)
    return buf.getvalue()
