"""API routes."""

from __future__ import annotations

import csv
import io

from fastapi import APIRouter, HTTPException, Query, UploadFile
from fastapi.responses import Response

import analytics
import reports
from db import connect, one, rows
from grading import grade_for, is_pass, points_for
from schemas import MarksBulkIn, ReportCard, ReportCardSubject, Student, StudentIn, Subject, SubjectIn

router = APIRouter(prefix="/api")


def get_db():
    conn = connect()
    try:
        yield conn
    finally:
        conn.close()


# ---------- students ----------

@router.get("/students", response_model=list[Student])
def list_students(class_name: str | None = None, q: str | None = None):
    conn = connect()
    try:
        sql = "SELECT * FROM students"
        where, params = [], []
        if class_name:
            where.append("class_name = ?")
            params.append(class_name)
        if q:
            where.append("(name LIKE ? OR admission_no LIKE ?)")
            params += [f"%{q}%", f"%{q}%"]
        if where:
            sql += " WHERE " + " AND ".join(where)
        sql += " ORDER BY class_name, name"
        return rows(conn.execute(sql, params))
    finally:
        conn.close()


@router.post("/students", response_model=Student, status_code=201)
def create_student(payload: StudentIn):
    conn = connect()
    try:
        cur = conn.execute(
            "INSERT INTO students (admission_no, name, class_name) VALUES (?, ?, ?)",
            (payload.admission_no, payload.name, payload.class_name))
        conn.commit()
        return one(conn.execute("SELECT * FROM students WHERE id = ?", (cur.lastrowid,)))
    except Exception:
        raise HTTPException(409, "Admission number already exists")
    finally:
        conn.close()


@router.put("/students/{student_id}", response_model=Student)
def update_student(student_id: int, payload: StudentIn):
    conn = connect()
    try:
        cur = conn.execute(
            "UPDATE students SET admission_no = ?, name = ?, class_name = ? WHERE id = ?",
            (payload.admission_no, payload.name, payload.class_name, student_id))
        conn.commit()
        if cur.rowcount == 0:
            raise HTTPException(404, "Student not found")
        return one(conn.execute("SELECT * FROM students WHERE id = ?", (student_id,)))
    finally:
        conn.close()


@router.delete("/students/{student_id}", status_code=204)
def delete_student(student_id: int):
    conn = connect()
    try:
        cur = conn.execute("DELETE FROM students WHERE id = ?", (student_id,))
        conn.commit()
        if cur.rowcount == 0:
            raise HTTPException(404, "Student not found")
    finally:
        conn.close()


@router.post("/students/import")
def import_students(file: UploadFile):
    """Import students from .xlsx or .csv (columns: admission_no, name, class_name)."""
    raw = file.file.read()
    records: list[tuple[str, str, str]] = []
    try:
        if file.filename and file.filename.lower().endswith(".csv"):
            text = raw.decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(text))
            records = [(r["admission_no"].strip(), r["name"].strip(), r["class_name"].strip())
                       for r in reader]
        else:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            ws = wb.active
            header = [str(c.value).strip().lower() if c.value else "" for c in ws[1]]
            idx = {name: header.index(name) for name in ("admission_no", "name", "class_name")
                   if name in header}
            if len(idx) < 3:
                raise ValueError("Header must include admission_no, name, class_name")
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row or not row[idx["admission_no"]]:
                    continue
                records.append((str(row[idx["admission_no"]]).strip(),
                                str(row[idx["name"]]).strip(),
                                str(row[idx["class_name"]]).strip()))
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(400, f"Could not parse file: {exc}")

    conn = connect()
    try:
        inserted = updated = 0
        for admission_no, name, class_name in records:
            if not admission_no or not name or not class_name:
                continue
            exists = conn.execute(
                "SELECT 1 FROM students WHERE admission_no = ?", (admission_no,)).fetchone()
            conn.execute(
                """INSERT INTO students (admission_no, name, class_name) VALUES (?, ?, ?)
                   ON CONFLICT(admission_no) DO UPDATE SET name = excluded.name,
                                                           class_name = excluded.class_name""",
                (admission_no, name, class_name))
            if exists:
                updated += 1
            else:
                inserted += 1
        conn.commit()
        return {"inserted": inserted, "updated": updated, "total": len(records)}
    finally:
        conn.close()


# ---------- subjects ----------

@router.get("/subjects", response_model=list[Subject])
def list_subjects():
    conn = connect()
    try:
        return rows(conn.execute("SELECT * FROM subjects ORDER BY code"))
    finally:
        conn.close()


@router.post("/subjects", response_model=Subject, status_code=201)
def create_subject(payload: SubjectIn):
    conn = connect()
    try:
        cur = conn.execute(
            "INSERT INTO subjects (name, code, teacher) VALUES (?, ?, ?)",
            (payload.name, payload.code, payload.teacher))
        conn.commit()
        return one(conn.execute("SELECT * FROM subjects WHERE id = ?", (cur.lastrowid,)))
    except Exception:
        raise HTTPException(409, "Subject name or code already exists")
    finally:
        conn.close()


# ---------- marks ----------

def _marks_sql(class_name: str | None, term: str | None, year: int | None) -> tuple[str, list]:
    sql = """SELECT m.student_id, s.name, s.admission_no, s.class_name,
                    sub.name AS subject, sub.code, m.score
             FROM marks m
             JOIN students s ON s.id = m.student_id
             JOIN subjects sub ON sub.id = m.subject_id"""
    where, params = [], []
    if class_name:
        where.append("s.class_name = ?")
        params.append(class_name)
    if term:
        where.append("m.term = ?")
        params.append(term)
    if year:
        where.append("m.year = ?")
        params.append(year)
    if where:
        sql += " WHERE " + " AND ".join(where)
    return sql, params


def fetch_marks(class_name: str | None = None, term: str | None = None,
                year: int | None = None) -> list[dict]:
    conn = connect()
    try:
        sql, params = _marks_sql(class_name, term, year)
        return rows(conn.execute(sql, params))
    finally:
        conn.close()


@router.get("/marks/matrix")
def marks_matrix(class_name: str = Query(...), term: str = Query(...), year: int = Query(...)):
    """Grid for marks entry: students x subjects for a class/term/year."""
    conn = connect()
    try:
        students = rows(conn.execute(
            "SELECT id, admission_no, name FROM students WHERE class_name = ? ORDER BY name",
            (class_name,)))
        subjects = rows(conn.execute("SELECT id, name, code FROM subjects ORDER BY code"))
        marks = rows(conn.execute(
            """SELECT student_id, subject_id, score FROM marks m
               JOIN students s ON s.id = m.student_id
               WHERE s.class_name = ? AND m.term = ? AND m.year = ?""",
            (class_name, term, year)))
        grid: dict[str, dict[str, float]] = {}
        for m in marks:
            grid.setdefault(str(m["student_id"]), {})[str(m["subject_id"])] = m["score"]
        return {"students": students, "subjects": subjects, "marks": grid}
    finally:
        conn.close()


@router.post("/marks/bulk")
def save_marks(payload: MarksBulkIn):
    conn = connect()
    try:
        cur = conn.cursor()
        for e in payload.entries:
            cur.execute(
                """INSERT INTO marks (student_id, subject_id, score, term, year)
                   VALUES (?, ?, ?, ?, ?)
                   ON CONFLICT(student_id, subject_id, term, year)
                   DO UPDATE SET score = excluded.score""",
                (e.student_id, e.subject_id, e.score, payload.term, payload.year))
        conn.commit()
        return {"saved": len(payload.entries)}
    finally:
        conn.close()


@router.delete("/marks/{mark_id}", status_code=204)
def delete_mark(mark_id: int):
    conn = connect()
    try:
        cur = conn.execute("DELETE FROM marks WHERE id = ?", (mark_id,))
        conn.commit()
        if cur.rowcount == 0:
            raise HTTPException(404, "Mark not found")
    finally:
        conn.close()


# ---------- dashboard & analytics ----------

@router.get("/dashboard")
def dashboard(class_name: str | None = None, term: str | None = None, year: int | None = None):
    marks = fetch_marks(class_name, term, year)
    overview = analytics.class_overview(marks)
    conn = connect()
    try:
        counts = {
            "students": conn.execute("SELECT COUNT(*) FROM students").fetchone()[0],
            "subjects": conn.execute("SELECT COUNT(*) FROM subjects").fetchone()[0],
            "classes": conn.execute(
                "SELECT COUNT(DISTINCT class_name) FROM students").fetchone()[0],
        }
    finally:
        conn.close()
    top = analytics.rankings(marks)[:5]
    return {**counts, **overview, "top_performers_list": top}


@router.get("/analytics/subjects")
def analytics_subjects(class_name: str | None = None, term: str | None = None,
                       year: int | None = None):
    data = analytics.subject_summary(fetch_marks(class_name, term, year))
    strongest = data[0] if data else None
    weakest = data[-1] if data else None
    return {"subjects": data, "strongest": strongest, "weakest": weakest}


@router.get("/analytics/rankings")
def analytics_rankings(class_name: str = Query(...), term: str | None = None,
                       year: int | None = None):
    return analytics.rankings(fetch_marks(class_name, term, year))


@router.get("/analytics/overview")
def analytics_overview(class_name: str | None = None, term: str | None = None,
                       year: int | None = None):
    return analytics.class_overview(fetch_marks(class_name, term, year))


# ---------- reports ----------

@router.get("/reports/student/{student_id}", response_model=ReportCard)
def student_report(student_id: int, term: str = Query(...), year: int = Query(...)):
    conn = connect()
    try:
        student = one(conn.execute("SELECT * FROM students WHERE id = ?", (student_id,)))
    finally:
        conn.close()
    if not student:
        raise HTTPException(404, "Student not found")

    rows_sub = analytics.student_report_rows(
        fetch_marks(student["class_name"], term, year), student_id)
    subjects = [ReportCardSubject(subject=r["subject"], code=r["code"], score=r["score"],
                                  grade=r["grade"], points=r["points"]) for r in rows_sub]
    if not subjects:
        raise HTTPException(404, "No marks recorded for this student in that term")
    scores = [s.score for s in subjects]
    average = round(sum(scores) / len(scores), 1)
    ranking = analytics.rankings(fetch_marks(student["class_name"], term, year))
    rank = next((r["rank"] for r in ranking if r["student_id"] == student_id), 0)
    from grading import aggregate_points, is_pass as _pass
    return ReportCard(
        student=Student(**student),
        term=term, year=year, subjects=subjects,
        total_points=aggregate_points(scores), average=average,
        mean_grade=grade_for(average), rank=rank, class_size=len(ranking),
        passed=_pass(average),
    )


@router.get("/reports/student/{student_id}/pdf")
def student_report_pdf(student_id: int, term: str = Query(...), year: int = Query(...)):
    report = student_report(student_id, term, year)
    pdf = reports.report_card_pdf(report.model_dump())
    return Response(pdf, media_type="application/pdf", headers={
        "Content-Disposition":
            f'attachment; filename="report-{report.student.admission_no}-{term.replace(" ", "")}{year}.pdf"'})


@router.get("/reports/class-ranking.xlsx")
def class_ranking_xlsx(class_name: str = Query(...), term: str = Query(...),
                       year: int = Query(...)):
    rows_ = analytics.rankings(fetch_marks(class_name, term, year))
    data = reports.class_ranking_xlsx(rows_, class_name, term, year)
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition":
                             f'attachment; filename="class-ranking-{term.replace(" ", "")}{year}.xlsx"'})


@router.get("/reports/subject-analysis.xlsx")
def subject_analysis_xlsx(term: str = Query(...), year: int = Query(...),
                          class_name: str | None = None):
    df = analytics.subject_analysis_frame(fetch_marks(class_name, term, year))
    data = reports.subject_analysis_xlsx(df, term, year)
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition":
                             f'attachment; filename="subject-analysis-{term.replace(" ", "")}{year}.xlsx"'})


# ---------- meta ----------

@router.get("/meta")
def meta():
    """Grading scale for the UI."""
    from grading import ORDER, POINTS
    return {
        "scale": [{"grade": g, "points": POINTS[g]} for g in ORDER],
        "pass_grade": "C+",
        "pass_mark": 55,
        "terms": ["Term 1", "Term 2", "Term 3"],
        "years": [2023, 2024, 2025],
        "classes": ["Form 1", "Form 2", "Form 3", "Form 4"],
    }
